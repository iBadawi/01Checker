#!/usr/bin/env python3
"""
build_pool.py — extract topics from both 01Gov Excel files and write pool.json

Sources:
  FILE1: 01Gov Content Topics - Sep 2025.xlsx
    - All "Rasha" sheets       → cols: country[1], agency[2], description[3], url[6]
    - Ibrahim Topics 2025      → cols: country[1], title[2], description[3], agency[5], url[6]
    - Ibrahim Topics 2026      → cols: country[1], title[2], description[3], agency[5], url[6]

  FILE2: _01GOV Portal Content Master Process 2026.xlsx
    - 2026 Content Dina        → cols: country[1], title[2], description[3], url[5]
"""

import json
import sys
from pathlib import Path
import openpyxl

FILE1 = Path("/Users/ibrhim/Downloads/01Gov Content Topics - Sep 2025.xlsx")
FILE2 = Path("/Users/ibrhim/Downloads/_01GOV Portal Content Master Process 2026.xlsx")
OUTPUT = Path(__file__).parent / "pool.json"


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def first_url(value) -> str:
    """Return the first http/https URL found in a cell (cells may contain multiple newline-separated URLs)."""
    for line in clean(value).splitlines():
        line = line.strip()
        if line.startswith("http://") or line.startswith("https://"):
            return line
    return ""


def find_sheet(wb, name: str):
    """Find sheet by stripped name (handles leading/trailing spaces in sheet names)."""
    return next((s for s in wb.sheetnames if s.strip() == name.strip()), None)


def extract_rasha(wb, sheet_name: str) -> list:
    """
    Rasha sheet layout (row 0 = header):
      [0]#  [1]country  [2]agency  [3]desc_en  [4]desc_ar  [5]category  [6]url  [7]aux_urls
    """
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    results = []
    for row in rows[1:]:  # skip header
        if not row or len(row) < 7:
            continue
        url = first_url(row[6])
        if not url:
            continue
        results.append({
            "source": sheet_name.strip(),
            "country": clean(row[1]),
            "agency": clean(row[2]),
            "practice_title": "",
            "description": clean(row[3]),
            "url": url,
        })
    return results


def extract_ibrahim(wb, sheet_name: str, label: str) -> list:
    """
    Ibrahim sheet layout:
      row 0 = section title (skip), row 1 = header (skip), row 2+ = data
      [0]#  [1]country  [2]title  [3]description  [4]domain  [5]agency  [6]url  [7]aux_urls
    """
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    results = []
    for row in rows[2:]:  # skip section title + header
        if not row or len(row) < 7:
            continue
        url = first_url(row[6])
        if not url:
            continue
        results.append({
            "source": label,
            "country": clean(row[1]),
            "agency": clean(row[5]),
            "practice_title": clean(row[2]),
            "description": clean(row[3]),
            "url": url,
        })
    return results


def extract_dina(wb, sheet_name: str) -> list:
    """
    Dina 2026 sheet layout (row 0 = header):
      [0]#  [1]country  [2]title  [3]description  [4]category  [5]url  [6]aux_urls  [7]week
    """
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    results = []
    for row in rows[1:]:  # skip header
        if not row or len(row) < 6:
            continue
        url = first_url(row[5])
        if not url:
            continue
        results.append({
            "source": "2026 Content Dina",
            "country": clean(row[1]),
            "agency": "",
            "practice_title": clean(row[2]),
            "description": clean(row[3]),
            "url": url,
        })
    return results


def main():
    pool = []
    seen_urls = set()

    def add(entries):
        for e in entries:
            if e["url"] and e["url"] not in seen_urls:
                seen_urls.add(e["url"])
                pool.append(e)

    # ── FILE 1 ──────────────────────────────────────────────────────────────
    print(f"Opening: {FILE1.name}")
    wb1 = openpyxl.load_workbook(FILE1, read_only=True, data_only=True)

    rasha_sheets = [s for s in wb1.sheetnames if "rasha" in s.strip().lower()]
    print(f"  Rasha sheets found: {len(rasha_sheets)}")
    for sheet_name in rasha_sheets:
        add(extract_rasha(wb1, sheet_name))

    for label in ("Ibrahim Topics 2025", "Ibrahim Topics 2026"):
        match = find_sheet(wb1, label)
        if match:
            entries = extract_ibrahim(wb1, match, label)
            print(f"  {label}: {len(entries)} topics")
            add(entries)
        else:
            print(f"  WARNING: sheet '{label}' not found", file=sys.stderr)

    wb1.close()

    # ── FILE 2 ──────────────────────────────────────────────────────────────
    print(f"Opening: {FILE2.name}")
    wb2 = openpyxl.load_workbook(FILE2, read_only=True, data_only=True)

    dina_match = find_sheet(wb2, "2026 Content Dina")
    if dina_match:
        entries = extract_dina(wb2, dina_match)
        print(f"  2026 Content Dina: {len(entries)} topics")
        add(entries)
    else:
        print("  WARNING: sheet '2026 Content Dina' not found", file=sys.stderr)

    wb2.close()

    # ── Output ───────────────────────────────────────────────────────────────
    with open(OUTPUT, "w", encoding="utf-8") as f:
        json.dump(pool, f, ensure_ascii=False, indent=2)

    print(f"\nDone — {len(pool)} unique topics written to {OUTPUT}")

    # Per-source summary
    from collections import Counter
    counts = Counter(e["source"] for e in pool)
    for source, count in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"  {count:4d}  {source}")


if __name__ == "__main__":
    main()
